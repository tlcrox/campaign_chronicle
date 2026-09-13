#!/usr/bin/env python3
"""merge.images.layout: how a scene image sits in the storyboard.

``chapter`` is the original document — one scene per page under a Heading 2.
``inline`` drops the headings and page breaks and sets a small picture in the
flow of dialogue, sized in pixels and aligned per config. Both are checked by
reading the saved .docx back, because the golden comparator sees only text: it
would notice a missing "Scene 01-001" heading and nothing else here.
"""

import base64
import csv
import tempfile
import unittest
from pathlib import Path
from unittest import mock

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

from pipeline.common.scenes import format_scene_name
from pipeline.config import Config
from pipeline.merge import storyboard
from pipeline.merge.storyboard import (EMU_PER_PX, Px, generate_storyboard,
                                       storyboard_filename)

# A real 1x1 JPEG: python-docx reads the header for the natural size.
JPG_1X1 = base64.b64decode(
    "/9j/4AAQSkZJRgABAQEASABIAAD/2wBDAAMCAgICAgMCAgIDAwMDBAYEBAQEBAgGBgUGCQgKCgkI"
    "CQkKDA8MCgsOCwkJDRENDg8QEBEQCgwSExIQEw8QEBD/yQALCAABAAEBAREA/8wABgAQEAX/2gAI"
    "AQEAAD8A0s8g/9k=")


def _images_block(layout="chapter", width_px=320, height_px=0, align="right"):
    return (
        "merge:\n"
        "  images:\n"
        "    width: 6.0\n"
        f"    layout: {layout}\n"
        "    inline:\n"
        f"      width_px: {width_px}\n"
        f"      height_px: {height_px}\n"
        f"      align: {align}\n"
    )


def _page_breaks(doc) -> int:
    return sum(1 for br in doc.element.body.iter(qn("w:br"))
               if br.get(qn("w:type")) == "page")


class _SessionCase(unittest.TestCase):
    """A one-video, two-scene session; dialogue before, between and after."""

    def setUp(self):
        self.t = tempfile.TemporaryDirectory()
        self.addCleanup(self.t.cleanup)
        self.root = Path(self.t.name)
        (self.root / "config").mkdir()
        self.images = self.root / "images"
        self.images.mkdir()
        for scene in (1, 2):
            (self.images / format_scene_name(1, scene, "jpg")).write_bytes(JPG_1X1)
        self.csv = self.root / "Scenes.csv"
        with open(self.csv, "w", encoding="utf-8", newline="") as fh:
            w = csv.writer(fh, lineterminator="\n")
            w.writerow(["Video", "Scene Number", "Start Time (seconds)"])
            w.writerow([1, 1, 10.0])
            w.writerow([1, 2, 20.0])
        self.transcript = self.root / "transcript.txt"
        self.transcript.write_text(
            "[0:05.00] Alice: before the first scene\n"
            "[0:15.00] Bob: between them\n"
            "[0:30.00] Alice: after the second\n",
            encoding="utf-8", newline="\n")
        self.out = self.root / "out.docx"

    def _config(self, **kw) -> Config:
        p = self.root / "config" / "config.yaml"
        p.write_text(_images_block(**kw), encoding="utf-8", newline="\n")
        return Config(p)

    def _build(self, **kw):
        cfg = self._config(**kw)
        with mock.patch.object(storyboard, "get_config", return_value=cfg):
            generate_storyboard(str(self.csv), str(self.transcript),
                                str(self.images), str(self.out))
        return Document(str(self.out))


class ChapterLayout(_SessionCase):
    """The default: unchanged, which is what keeps the goldens passing."""

    def test_default_layout_is_chapter(self):
        self.assertEqual(self._config().storyboard_layout, "chapter")

    def test_one_heading_and_one_page_break_per_scene(self):
        doc = self._build()
        headings = [p.text for p in doc.paragraphs if p.style.name == "Heading 2"]
        self.assertEqual(headings, ["Scene 01-001", "Scene 01-002"])
        # Title page break + one before each scene that follows dialogue.
        self.assertEqual(_page_breaks(doc), 3)
        self.assertEqual(len(doc.inline_shapes), 2)


class InlineLayout(_SessionCase):
    def test_no_headings_and_no_scene_page_breaks(self):
        doc = self._build(layout="inline")
        self.assertEqual([p.text for p in doc.paragraphs if p.style.name.startswith("Heading")],
                         [])
        self.assertEqual(_page_breaks(doc), 1, "only the title page's break")
        self.assertEqual(len(doc.inline_shapes), 2)

    def test_dialogue_is_the_same_in_both_layouts(self):
        chapter = [p.text for p in self._build().paragraphs
                   if p.text and not p.text.startswith("Scene ")]
        inline = [p.text for p in self._build(layout="inline").paragraphs if p.text]
        self.assertEqual(chapter, inline)

    def test_pictures_sit_between_the_dialogue_they_precede(self):
        doc = self._build(layout="inline")
        kinds = []
        for p in doc.paragraphs:
            if p._element.findall(".//" + qn("w:drawing")):
                kinds.append("picture")
            elif p.text:
                kinds.append(p.text.split(":")[0])
        self.assertEqual(kinds, ["Transcript",  # the document title
                                 "Alice", "picture", "Bob", "picture", "Alice"])

    def test_width_only_keeps_the_aspect_ratio(self):
        doc = self._build(layout="inline", width_px=320, height_px=0)
        for shape in doc.inline_shapes:
            self.assertEqual(shape.width, 320 * EMU_PER_PX)
            self.assertEqual(shape.height, 320 * EMU_PER_PX, "1x1 source stays square")

    def test_both_dimensions_scale_to_that_box(self):
        doc = self._build(layout="inline", width_px=320, height_px=180)
        for shape in doc.inline_shapes:
            self.assertEqual((shape.width, shape.height),
                             (Px(320), Px(180)))

    def test_alignment_is_applied_to_the_picture_paragraph(self):
        expected = {"left": WD_ALIGN_PARAGRAPH.LEFT,
                    "center": WD_ALIGN_PARAGRAPH.CENTER,
                    "right": WD_ALIGN_PARAGRAPH.RIGHT}
        for align, want in expected.items():
            with self.subTest(align=align):
                doc = self._build(layout="inline", align=align)
                picture_paragraphs = [p for p in doc.paragraphs
                                      if p._element.findall(".//" + qn("w:drawing"))]
                self.assertEqual(len(picture_paragraphs), 2)
                for p in picture_paragraphs:
                    self.assertEqual(p.alignment, want)


class LayoutConfigIsValidatedAtLoad(_SessionCase):
    """Bad values fail with the rest of the report, not after a two-hour run."""

    def _problems(self, **kw) -> str:
        with self.assertRaises(ValueError) as cm:
            self._config(**kw)
        return str(cm.exception)

    def test_unknown_layout(self):
        self.assertIn("merge.images.layout must be one of: chapter, inline",
                      self._problems(layout="floating"))

    def test_unknown_alignment(self):
        self.assertIn("merge.images.inline.align must be one of: left, center, right",
                      self._problems(align="justify"))

    def test_size_must_be_whole_pixels(self):
        self.assertIn("merge.images.inline.width_px must be a whole number",
                      self._problems(width_px="wide"))

    def test_inline_needs_at_least_one_dimension(self):
        self.assertIn("needs width_px or height_px",
                      self._problems(layout="inline", width_px=0, height_px=0))

    def test_chapter_does_not_care_about_inline_size(self):
        cfg = self._config(layout="chapter", width_px=0, height_px=0)
        self.assertEqual(cfg.storyboard_layout, "chapter")

    def test_all_problems_are_reported_together(self):
        report = self._problems(layout="floating", align="justify", width_px=-1)
        for fragment in ("layout must be", "align must be", "width_px must be"):
            self.assertIn(fragment, report)


class LayoutOverride(_SessionCase):
    """``layout=`` on the call beats the config, for one document only."""

    def _build_with(self, layout_arg, **cfg_kw):
        cfg = self._config(**cfg_kw)
        with mock.patch.object(storyboard, "get_config", return_value=cfg):
            generate_storyboard(str(self.csv), str(self.transcript),
                                str(self.images), str(self.out), layout=layout_arg)
        return Document(str(self.out))

    def test_inline_argument_beats_chapter_config(self):
        doc = self._build_with("inline", layout="chapter")
        self.assertEqual(_page_breaks(doc), 1)
        self.assertEqual([p.text for p in doc.paragraphs if p.style.name == "Heading 2"], [])

    def test_chapter_argument_beats_inline_config(self):
        doc = self._build_with("chapter", layout="inline")
        self.assertEqual([p.text for p in doc.paragraphs if p.style.name == "Heading 2"],
                         ["Scene 01-001", "Scene 01-002"])

    def test_none_means_the_config(self):
        doc = self._build_with(None, layout="inline")
        self.assertEqual(_page_breaks(doc), 1)


class StageNamesTheDocumentForItsLayout(_SessionCase):
    """cc_stages.generate_storyboard: <session>_storyboard.docx or
    <session>_inline.docx, so both can live in one cc_output/."""

    def test_filename_rule(self):
        self.assertEqual(storyboard_filename("Week 13", "chapter"), "Week 13_storyboard.docx")
        self.assertEqual(storyboard_filename("Week 13", "inline"), "Week 13_inline.docx")

    def test_stage_default_output_follows_the_effective_layout(self):
        from cc_stages import generate_storyboard as stage
        cfg = self._config(layout="chapter")
        session = self.root / "Week 13"
        combined = session / "cc_output" / "combined_output"
        combined.mkdir(parents=True)
        (combined / "Week 13_transcript_combined.txt").write_text(
            self.transcript.read_text(encoding="utf-8"), encoding="utf-8")
        with mock.patch.object(stage, "get_config", return_value=cfg), \
                mock.patch.object(storyboard, "get_config", return_value=cfg):
            self.assertTrue(stage.run(session_dir=session, config=cfg))
            self.assertTrue(stage.run(session_dir=session, config=cfg, layout="inline"))
        names = sorted(p.name for p in (session / "cc_output").glob("*.docx"))
        self.assertEqual(names, ["Week 13_inline.docx", "Week 13_storyboard.docx"])


class PixelUnit(unittest.TestCase):
    def test_px_is_word_96_dpi(self):
        self.assertEqual(Px(96), 914400, "96 px is one inch in EMU")


if __name__ == "__main__":
    unittest.main()
